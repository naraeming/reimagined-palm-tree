"""Independent verification of original exports, aggregate counts and XLSX output."""
import argparse
import hashlib
import json
import pathlib
from openpyxl import load_workbook

parser=argparse.ArgumentParser()
parser.add_argument('--input',type=pathlib.Path,required=True)
parser.add_argument('--report-dir',type=pathlib.Path,required=True)
args=parser.parse_args()
d=json.loads((args.report_dir/'aggregates/analysis.json').read_text(encoding='utf-8'))
independent_total=0
for source in d['sources']:
    path=args.input/source['file']
    assert hashlib.sha256(path.read_bytes()).hexdigest()==source['sha256'], 'Source was changed'
    book=load_workbook(path,read_only=True,data_only=True)
    rows=list(book.active.iter_rows(min_row=2,values_only=True))
    rows=[row for row in rows if any(v is not None for v in row)]
    unique=set(rows)
    assert len(rows)==source['raw_rows']
    assert len(unique)==source['accounts'], 'Whole-row uniqueness differs from account grouping'
    independent_total+=len(unique)
    book.close()
assert independent_total==d['totals']['accounts']
for dimension in ['languages','countries','cities','country_language','monthly_language','monthly_country','order_buckets']:
    assert sum(row['accounts'] for row in d[dimension])==independent_total, dimension
if not d['scope']['growth_available']:
    assert all(row['mom_growth'] is None and row['yoy_growth'] is None for row in d['monthly'])
    assert all(row['full_year_growth'] is None for row in d['annual'])
    assert all(row['same_period_growth'] is None for row in d['matched_periods'])
report_name=f"회원분석_{d['scope']['first_month'].replace('-','')}-{d['scope']['last_month'].replace('-','')}_제공분검증.xlsx"
path=args.report_dir/report_name
book=load_workbook(path,data_only=True)
assert book['요약']['B10'].value==independent_total
assert book['요약']['B8'].value==d['totals']['raw_rows']
assert book['요약']['B9'].value==d['totals']['duplicate_rows_removed']
errors={'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}
for sheet in book:
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type != 'e', f'Error in {sheet.title}:{cell.coordinate}'
            if isinstance(cell.value,str):
                assert cell.value not in errors, f'Error in {sheet.title}:{cell.coordinate}'
if not d['scope']['growth_available']:
    for row in range(7,len(d['monthly'])+7):
        assert book['월별집계'][f'L{row}'].value=='n.a.'
        assert book['월별집계'][f'M{row}'].value=='n.a.'
assert len(book['요약']._charts)==1
assert len(book['언어별']._charts)==1
book.close()
result={'source_files_unchanged':len(d['sources']),'independent_unique_rows':independent_total,
        'aggregate_dimensions_reconciled':7,'workbook_formula_errors':0,'native_charts':2,
        'growth_available':d['scope']['growth_available'],'verification':'passed'}
(args.report_dir/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
