"""Independently reconcile an exported monthly order workbook to its source."""
import argparse
import calendar
from collections import defaultdict
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import unicodedata

from openpyxl import load_workbook


def decimal(value):
    return Decimal(str(value).replace(',', '').strip())


def close(actual, expected, tolerance=Decimal('0.0001')):
    assert actual is not None, 'Missing cached formula value'
    assert abs(decimal(actual) - decimal(expected)) <= tolerance, (actual, expected)


def text(value, upper=False):
    if value is None or not str(value).strip():
        return '미기록'
    value = unicodedata.normalize('NFC', str(value).strip())
    return value.upper() if upper else value.lower()


def verify(source, report_dir):
    if isinstance(source,list):
        if len(source)>1:return verify_trends(source,report_dir)
        source=source[0]
    d = json.loads((report_dir / 'aggregates/order-analysis.json').read_text(encoding='utf-8'))
    assert len(d['sources']) == 1, 'Verifier expects one monthly source'
    assert source.name == d['sources'][0]['file']
    assert hashlib.sha256(source.read_bytes()).hexdigest() == d['sources'][0]['sha256'], 'Source changed'
    expected = []
    src = load_workbook(source, read_only=True, data_only=True)
    for row_number, values in enumerate(src.active.iter_rows(values_only=True), 1):
        if all(v is None for v in values):
            continue
        if row_number==1 and d['sources'][0].get('header_present'):
            assert list(values)==(['통화','국가','언어','주문상태','업종','상품금액','최종금액','주문수'] if len(values)==8 else ['통화','국가','도시','업종','주문언어','주문상태','상품금액','최종금액','주문수'])
            continue
        if len(values)==8:
            values=(values[0],values[1],'도시 항목 없음',values[4],values[2],values[3],values[5],values[6],values[7])
        assert len(values) == 9
        status = 'blank' if values[5] is None or str(values[5]).strip() == '' else str(int(decimal(values[5])))
        status_meta = d['scope']['status_mapping'].get(status, {'label': f'상태 {status}(정의 미확인)', 'group': 'unknown'})
        expected.append({
            'month': f'20{source.stem[:2]}-{source.stem[2:]}',
            **dict(zip(['currency', 'country', 'city', 'industry', 'language'],
                       [text(v, i < 2) for i, v in enumerate(values[:5])])),
            'status': status, 'status_label': status_meta['label'], 'status_group': status_meta['group'],
            'product_amount': decimal(values[6]), 'final_amount': decimal(values[7]),
            'orders': int(decimal(values[8])), 'source_row': row_number,
        })
    src.close()
    assert len(expected) == len(d['rows']) == d['scope']['aggregate_rows']
    for original, normalized in zip(expected, d['rows']):
        for field, value in original.items():
            assert value == (decimal(normalized[field]) if field.endswith('_amount') else normalized[field]), field

    partitions = {
        'monthly': ['month'], 'status': ['status', 'status_label', 'status_group'],
        'countries': ['country'], 'cities': ['country', 'city'], 'languages': ['language'],
        'industries': ['industry'], 'country_language': ['country', 'language'],
        'country_industry': ['country', 'industry'], 'currency': ['currency'],
        'monthly_country': ['month', 'country'], 'monthly_language': ['month', 'language'],
    }
    for table, dimensions in partitions.items():
        independent = defaultdict(lambda: defaultdict(int))
        for r in expected:
            key = tuple(r[k] for k in dimensions)
            independent[key]['orders'] += r['orders']
            independent[key][r['status_group']] += r['orders']
        assert len(independent) == len(d[table]), table
        for group in d[table]:
            key = tuple(group[k] for k in dimensions)
            assert group['orders'] == independent[key]['orders'], table
            for status_group, count_field in [('completed', 'completed_orders'), ('cancelled', 'cancelled_orders'),
                                               ('known_other', 'known_other_orders'), ('unknown', 'unknown_status_orders')]:
                assert group[count_field] == independent[key][status_group], table
    total = sum(r['orders'] for r in expected)
    assert total == d['totals']['orders']
    name = f"주문집계_20{source.stem}_분석.xlsx"
    report = report_dir / name
    values = load_workbook(report, data_only=True)
    formulas = load_workbook(report, data_only=False)
    raw = values['집계자료']
    assert raw.max_row == len(expected) + 5
    for n, r in enumerate(expected, 6):
        wanted = [r['month'], r['currency'], r['country'], r['city'], r['industry'], r['language'],
                  '공란' if r['status'] == 'blank' else r['status'], r['status_label'], r['status_group'],
                  r['product_amount'], r['final_amount'], r['orders'], r['source_row']]
        for col, item in enumerate(wanted, 1):
            actual = raw.cell(n, col).value
            if isinstance(item, Decimal):
                close(actual, item)
            else:
                assert actual == item, (n, col)
    close(values['주문요약'].cell(7 + len(d['status']), 2).value, total)
    for i, r in enumerate(d['status'], 7):
        assert values['주문요약'].cell(i, 1).value == r['status_label']
        close(values['주문요약'].cell(i, 2).value, r['orders'])
    for sheet_name, start, table in [('국가도시', 7, 'countries'), ('언어업종', 7, 'languages'), ('언어업종', 18, 'industries')]:
        for i, r in enumerate(d[table], start):
            for col, count in [(2, r['orders']), (3, r['completed_orders']), (4, r['cancelled_orders']),
                               (5, r['known_other_orders'] + r['unknown_status_orders'])]:
                close(values[sheet_name].cell(i, col).value, count)
            close(values[sheet_name].cell(i, 6).value, r['orders'] / total)
            close(values[sheet_name].cell(i, 7).value, r['cancellation_rate'])
    for i, r in enumerate(d['cities'], 21):
        close(values['국가도시'].cell(i, 3).value, r['orders'])
        close(values['국가도시'].cell(i, 4).value, r['completed_orders'])
        close(values['국가도시'].cell(i, 5).value, r['cancelled_orders'])
    for i, country in enumerate(d['countries'], 7):
        for j, language in enumerate(d['languages'], 2):
            count = sum(r['orders'] for r in expected if r['country'] == country['country'] and r['language'] == language['language'])
            close(values['국가언어'].cell(i, j).value, count)
    for i, r in enumerate(d['currency'], 7):
        completed = [x for x in expected if x['currency'] == r['currency'] and x['status_group'] == 'completed']
        product = sum((x['product_amount'] for x in completed), Decimal(0))
        final = sum((x['final_amount'] for x in completed), Decimal(0))
        assert product == decimal(r['completed_product_amount'])
        assert final == decimal(r['completed_final_amount'])
        close(values['통화금액'].cell(i, 2).value, r['completed_orders'])
        close(values['통화금액'].cell(i, 3).value, product)
        close(values['통화금액'].cell(i, 4).value, final)
        if d['scope']['amounts_are_sums_confirmed'] and completed:
            close(values['통화금액'].cell(i, 5).value, final / r['completed_orders'])
        close(values['통화금액'].cell(i, 6).value, final - product)
    formula_count = 0
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                cached = values[sheet.title][cell.coordinate]
                assert cached.data_type != 'e', (sheet.title, cell.coordinate, cached.value)
                if cell.data_type == 'f':
                    formula_count += 1
                    assert cached.value is not None, (sheet.title, cell.coordinate)
    assert len(formulas['주문요약']._charts) == 1, 'Missing country chart'
    result = {'passed': True, 'source_rows': len(expected), 'orders': total, 'formula_cells': formula_count,
              'source_hash_matches': True, 'first_data_row_included': True, 'partitions_reconciled': len(partitions),
              'currency_amounts_reconciled': True, 'admin_total_matches': d['scope']['admin_total_matches']}
    values.close()
    formulas.close()
    (report_dir / 'order-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False))


def verify_trends(sources, report_dir):
    d=json.loads((report_dir/'aggregates/order-analysis.json').read_text(encoding='utf-8'))
    layout=json.loads((report_dir/'order-layout.json').read_text(encoding='utf-8'))
    actual_sources={p.name:p for p in sources}
    assert len(actual_sources)==len(sources)==len(d['sources'])
    original=[]
    for meta in d['sources']:
        source=actual_sources[meta['file']]
        assert hashlib.sha256(source.read_bytes()).hexdigest()==meta['sha256'],source.name
        wb=load_workbook(source,read_only=True,data_only=True)
        loaded=0
        for n,row in enumerate(wb.active.iter_rows(values_only=True),1):
            if all(x is None for x in row):continue
            if n==1 and meta['header_present']:
                assert list(row)==(['통화','국가','언어','주문상태','업종','상품금액','최종금액','주문수'] if len(row)==8 else ['통화','국가','도시','업종','주문언어','주문상태','상품금액','최종금액','주문수'])
                continue
            if len(row)==8:
                currency,country,language,status,industry,product,final,count=row
                city='도시 항목 없음'
            else:
                assert len(row)==9
                currency,country,city,industry,language,status,product,final,count=row
            code='blank' if status is None or str(status).strip()=='' else str(int(decimal(status)))
            sm=d['scope']['status_mapping'].get(code,{'label':f'상태 {code}(정의 미확인)','group':'unknown'})
            original.append({'month':meta['month'],'currency':text(currency,True),'country':text(country,True),
                'city':text(city),'industry':text(industry),'language':text(language),'status':code,'status_label':sm['label'],
                'status_group':sm['group'],'product_amount':decimal(product),'final_amount':decimal(final),
                'orders':int(decimal(count)),'source_row':n,'source_file':source.name,'city_available':len(row)==9})
            loaded+=1
        assert loaded==meta['rows']
        wb.close()
    assert len(original)==len(d['rows'])
    for r,expected in zip(original,d['rows']):
        for field,value in r.items():
            assert value==(decimal(expected[field]) if field.endswith('_amount') else expected[field]),field
    def filtered(**criteria):return [r for r in original if all(r[k]==v for k,v in criteria.items())]
    def orders(**criteria):return sum(r['orders'] for r in filtered(**criteria))
    tables={'monthly':['month'],'status':['status'],'countries':['country'],'cities':['country','city'],
            'languages':['language'],'industries':['industry'],'country_language':['country','language'],
            'country_industry':['country','industry'],'currency':['currency'],'monthly_country':['month','country'],
            'monthly_language':['month','language'],'monthly_industry':['month','industry'],'monthly_currency':['month','currency']}
    for table,keys in tables.items():
        assert sum(r['orders'] for r in d[table])==orders(),table
        for group in d[table]:
            criteria={k:group[k] for k in keys}
            assert group['orders']==orders(**criteria),table
            for status,countkey in [('completed','completed_orders'),('cancelled','cancelled_orders'),('known_other','known_other_orders'),('unknown','unknown_status_orders')]:
                assert group[countkey]==orders(**criteria,status_group=status),table
    assert {r['month'] for r in d['monthly_city']}==set(d['scope']['city_months'])
    for group in d['monthly_city']:
        assert group['orders']==orders(month=group['month'],country=group['country'],city=group['city'])
    months=d['scope']['months']
    pending=set(d['scope'].get('replacement_pending_months',[]))
    report=report_dir/f"주문추이_{months[0].replace('-','')}_{months[-1].replace('-','')}.xlsx"
    values=load_workbook(report,data_only=True);formulas=load_workbook(report,data_only=False)
    raw=values['집계자료'];assert raw.max_row==len(original)+5
    for row,r in enumerate(original,6):
        wanted=[r['month'],r['currency'],r['country'],r['city'],r['industry'],r['language'],
                '공란' if r['status']=='blank' else r['status'],r['status_label'],r['status_group'],
                r['product_amount'],r['final_amount'],r['orders'],r['source_row'],None,r['source_file']]
        for c,w in enumerate(wanted,1):
            if isinstance(w,Decimal):close(raw.cell(row,c).value,w)
            else:assert raw.cell(row,c).value==w,(row,c)
    for row,m in enumerate(months,6):
        sh=values['월별추이'];assert sh.cell(row,1).value==m
        count=orders(month=m)
        for c,status in enumerate([None,'completed','cancelled','known_other','unknown'],2):
            close(sh.cell(row,c).value,count if status is None else orders(month=m,status_group=status))
        year,month=map(int,m.split('-'));days=calendar.monthrange(year,month)[1]
        close(sh.cell(row,7).value,orders(month=m,status_group='cancelled')/count)
        close(sh.cell(row,8).value,days);close(sh.cell(row,9).value,count/days)
        prev=f'{year-1}-12' if month==1 else f'{year}-{month-1:02}'
        for c,prior,daily in [(10,prev,False),(11,prev,True),(12,f'{year-1}-{month:02}',False)]:
            if prior not in months or m in pending or prior in pending:assert sh.cell(row,c).value=='n.a.'
            else:
                expected=count/orders(month=prior)-1
                if daily:expected=(count/days)/(orders(month=prior)/calendar.monthrange(*map(int,prior.split('-')))[1])-1
                close(sh.cell(row,c).value,expected)
        assert sh.cell(row,13).value==('교체 예정' if m in pending else '제공본')
    total_row=len(months)+6
    close(values['월별추이'].cell(total_row,2).value,orders())
    close(values['월별추이'].cell(total_row,7).value,orders(status_group='cancelled')/orders())
    close(values['월별추이'].cell(total_row,9).value,orders()/sum(calendar.monthrange(*map(int,m.split('-')))[1] for m in months))
    for block in layout['matrix']:
        sh=values['분류별추이']
        for row,group in enumerate(d[block['table']],block['start']):
            key=group[block['key']];assert sh.cell(row,1).value==key
            for c,m in enumerate(months,2):close(sh.cell(row,c).value,orders(month=m,**{block['key']:key}))
            close(sh.cell(row,len(months)+2).value,orders(**{block['key']:key}))
            close(sh.cell(row,len(months)+3).value,orders(month=months[-1],**{block['key']:key})/orders(month=months[-1]))
    for row,group in enumerate(layout['cities'],6):
        sh=values['도시추이'];assert [sh.cell(row,c).value for c in (1,2)]==[group['country'],group['city']]
        for c,m in enumerate(months,3):
            if m not in d['scope']['city_months']:assert sh.cell(row,c).value=='n.a.'
            else:close(sh.cell(row,c).value,orders(month=m,**group))
    geographic_growth_rows=0
    for block in layout.get('geography_growth',[]):
        sh=values[block['sheet']];is_city=block['is_city'];first_value=4 if is_city else 3
        groups=layout['cities'] if is_city else [{'country':r['country']} for r in d['countries']]
        available=lambda m:m in months and (not is_city or m in d['scope']['city_months'])
        for i,group in enumerate(groups):
            for j,m in enumerate(months):
                row=block['start']+i*len(months)+j;geographic_growth_rows+=1
                labels=[group['country']]+([group['city']] if is_city else [])+[m]
                assert [sh.cell(row,c).value for c in range(1,first_value)]==labels
                year,month=map(int,m.split('-'));prev=f'{year-1}-12' if month==1 else f'{year}-{month-1:02}';py=f'{year-1}-{month:02}'
                current=orders(month=m,**group)
                for offset,period in [(0,m),(1,prev),(5,py)]:
                    actual=sh.cell(row,first_value+offset).value
                    if available(period):close(actual,orders(month=period,**group))
                    else:assert actual=='n.a.'
                for previous,delta_offset,rate_offset in [(prev,2,3),(py,6,7)]:
                    ready=available(m) and available(previous) and m not in pending and previous not in pending
                    base=orders(month=previous,**group) if available(previous) else 0
                    if ready:
                        close(sh.cell(row,first_value+delta_offset).value,current-base)
                        if base:close(sh.cell(row,first_value+rate_offset).value,current/base-1)
                        else:assert sh.cell(row,first_value+rate_offset).value=='n.a.'
                    else:
                        assert sh.cell(row,first_value+delta_offset).value=='n.a.'
                        assert sh.cell(row,first_value+rate_offset).value=='n.a.'
                base=orders(month=prev,**group) if available(prev) else 0
                if available(m) and available(prev) and base and m not in pending and prev not in pending:
                    daily=(current/calendar.monthrange(year,month)[1])/(base/calendar.monthrange(*map(int,prev.split('-')))[1])-1
                    close(sh.cell(row,first_value+4).value,daily)
                else:assert sh.cell(row,first_value+4).value=='n.a.'
    for row,group in enumerate(layout['currency'],6):
        sh=values['통화별추이'];rows=filtered(**group,status_group='completed')
        assert [sh.cell(row,c).value for c in (1,2)]==[group['month'],group['currency']]
        count=sum(r['orders'] for r in rows)
        product=sum((r['product_amount'] for r in rows),Decimal(0));final=sum((r['final_amount'] for r in rows),Decimal(0))
        close(sh.cell(row,3).value,count);close(sh.cell(row,4).value,product);close(sh.cell(row,5).value,final)
        if count and d['scope']['amounts_are_sums_confirmed']:close(sh.cell(row,6).value,final/count)
        close(sh.cell(row,9).value,orders(**group))
        year,month=map(int,group['month'].split('-'));prev=f'{year-1}-12' if month==1 else f'{year}-{month-1:02}'
        prev_rows=filtered(month=prev,currency=group['currency'],status_group='completed')
        prev_final=sum((r['final_amount'] for r in prev_rows),Decimal(0));prev_count=sum(r['orders'] for r in prev_rows)
        ready=group['month'] not in pending and prev not in pending
        if prev_final and ready:close(sh.cell(row,7).value,final/prev_final-1)
        else:assert sh.cell(row,7).value=='n.a.'
        if prev_final and prev_count and count and d['scope']['amounts_are_sums_confirmed'] and ready:close(sh.cell(row,8).value,(final/count)/(prev_final/prev_count)-1)
        else:assert sh.cell(row,8).value=='n.a.'
    formula_count=0
    for sh in formulas:
        for row in sh:
            for cell in row:
                cached=values[sh.title][cell.coordinate]
                assert cached.data_type!='e',(sh.title,cell.coordinate,cached.value)
                if cell.data_type=='f':
                    formula_count+=1
                    assert cached.value is not None,(sh.title,cell.coordinate)
    assert len(formulas['월별추이']._charts)==1
    result={'passed':True,'months':months,'source_rows':len(original),'orders':orders(),'formula_cells':formula_count,
            'source_hashes_match':True,'headers_and_first_data_rows_verified':True,'city_absence_not_zero':True,
            'monthly_and_currency_growth_verified':True,'geographic_growth_rows_verified':geographic_growth_rows,
            'replacement_pending_months':sorted(pending),'admin_total_matches':d['scope']['admin_total_matches']}
    values.close();formulas.close()
    (report_dir/'order-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path, nargs='+', required=True)
    parser.add_argument('--report-dir', type=Path, required=True)
    args = parser.parse_args()
    verify(args.input, args.report_dir)
