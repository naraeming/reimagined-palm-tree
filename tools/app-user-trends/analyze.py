"""Audit monthly registration exports and emit aggregate-only analysis.

No names, emails, IP addresses or account identifiers are written to outputs.
Run with --help. A YYMM filename denotes the registration month, not a snapshot month.
"""
from __future__ import annotations

import argparse
import collections
import csv
import datetime as dt
import hashlib
import json
import math
import pathlib
import re
import unicodedata

from openpyxl import load_workbook


def month_key(stem):
    if not re.fullmatch(r"\d{4}", stem):
        raise ValueError("Expected YYMM filename")
    year, month = 2000 + int(stem[:2]), int(stem[2:])
    dt.date(year, month, 1)
    return f"{year:04}-{month:02}"


def account_id(value):
    text = str(value or "")
    match = re.match(r"^\s*(\d+)(?=\s*(?:\.|$))", text)
    return match.group(1) if match else None


def category(value, upper=False):
    if value is None or str(value).strip() == "":
        return "미기록"
    text = unicodedata.normalize("NFC", str(value).strip())
    return text.upper() if upper else text.lower()


def numeric(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        number = float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return "invalid"
    return int(number) if math.isfinite(number) and number >= 0 and number.is_integer() else "invalid"


def rates(current, previous):
    return None if previous in (None, 0) else current / previous - 1


def csv_out(path, rows):
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def analyze(input_dir, output_dir, start=None, end=None, completeness_confirmed=False, blanks_are_zero=False, reference_counts=None):
    files = []
    for path in input_dir.glob("*.xlsx"):
        if not re.fullmatch(r"\d{4}", path.stem):
            continue
        month = month_key(path.stem)
        if (start is None or month >= start) and (end is None or month <= end):
            files.append((month, path))
    files.sort()
    if not files:
        raise ValueError("No eligible monthly files")
    all_accounts = {}
    manifests, monthly = [], []
    language = collections.Counter()
    country = collections.Counter()
    city = collections.Counter()
    country_language = collections.Counter()
    monthly_language = collections.Counter()
    monthly_country = collections.Counter()
    level = collections.Counter()
    buckets = collections.Counter()
    issues = []
    for month, path in files:
        wb = load_workbook(path, read_only=True, data_only=True)
        if len(wb.worksheets) != 1:
            raise ValueError(f"Unexpected sheet count in {path.name}")
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        header = list(next(iterator))
        required = {"id", "locale", "country", "city", "orders", "reviews", "shops", "level", "disabled"}
        if not required.issubset(header):
            raise ValueError(f"Missing required columns in {path.name}")
        seen = {}
        raw_rows = duplicate_rows = conflicting_ids = invalid_ids = 0
        for row in iterator:
            if all(value is None for value in row):
                continue
            raw_rows += 1
            data = dict(zip(header, row))
            aid = account_id(data["id"])
            if not aid:
                invalid_ids += 1
                continue
            if aid in seen:
                duplicate_rows += 1
                if seen[aid] != data:
                    conflicting_ids += 1
                continue
            seen[aid] = data
        wb.close()
        if invalid_ids or conflicting_ids:
            raise ValueError(f"Cannot safely deduplicate {path.name}: invalid_ids={invalid_ids}, conflicts={conflicting_ids}")
        summary = {"month": month, "file": path.name, "raw_rows": raw_rows, "accounts": len(seen),
                   "duplicate_rows_removed": duplicate_rows, "orders_recorded_accounts": 0,
                   "orders_positive_accounts": 0, "orders_blank_accounts": 0, "orders_invalid_accounts": 0,
                   "recorded_cumulative_orders": 0, "reviews_positive_accounts": 0,
                   "reviews_blank_accounts": 0, "shops_field_present_accounts": 0,
                   "disabled_true_accounts": 0, "disabled_blank_accounts": 0,
                   "country_missing_accounts": 0, "city_missing_accounts": 0,
                   "locale_missing_accounts": 0, "phone_verification_present_accounts": 0}
        for aid, data in seen.items():
            if aid in all_accounts:
                raise ValueError(f"Cross-month account overlap between {all_accounts[aid]} and {month}; review cohort filters")
            all_accounts[aid] = month
            lang = category(data["locale"])
            nation = category(data["country"], upper=True)
            town = category(data["city"])
            language[lang] += 1
            country[nation] += 1
            city[(nation, town)] += 1
            country_language[(nation, lang)] += 1
            monthly_language[(month, lang)] += 1
            monthly_country[(month, nation)] += 1
            level[category(data["level"])] += 1
            summary["locale_missing_accounts"] += lang == "미기록"
            summary["country_missing_accounts"] += nation == "미기록"
            summary["city_missing_accounts"] += town == "미기록"
            order_count = numeric(data["orders"])
            if order_count is None:
                summary["orders_blank_accounts"] += 1
                bucket = "0건(빈칸 포함)" if blanks_are_zero else "빈칸(의미 미확인)"
            elif order_count == "invalid":
                summary["orders_invalid_accounts"] += 1
                bucket = "비정상 값"
            else:
                summary["orders_recorded_accounts"] += 1
                summary["orders_positive_accounts"] += order_count > 0
                summary["recorded_cumulative_orders"] += order_count
                bucket = "0건(빈칸 포함)" if order_count == 0 and blanks_are_zero else ("0건" if order_count == 0 else "1건" if order_count == 1 else "2~4건" if order_count <= 4 else "5~9건" if order_count <= 9 else "10건 이상")
            buckets[(month, bucket)] += 1
            review_count = numeric(data["reviews"])
            summary["reviews_positive_accounts"] += isinstance(review_count, int) and review_count > 0
            summary["reviews_blank_accounts"] += review_count is None
            summary["shops_field_present_accounts"] += data["shops"] is not None and str(data["shops"]).strip() != ""
            summary["disabled_true_accounts"] += data["disabled"] is True or str(data["disabled"]).lower() == "true"
            summary["disabled_blank_accounts"] += data["disabled"] is None or data["disabled"] == ""
            summary["phone_verification_present_accounts"] += data.get("telephone_verified_at") is not None and data.get("telephone_verified_at") != ""
        monthly.append(summary)
        manifests.append({"file":path.name,"month":month,"bytes":path.stat().st_size,
                          "sha256":hashlib.sha256(path.read_bytes()).hexdigest(),"sheet":ws.title,
                          "raw_rows":raw_rows,"accounts":len(seen),"duplicate_rows_removed":duplicate_rows,
                          "conflicting_duplicate_rows":conflicting_ids,"invalid_id_rows":invalid_ids,
                          "file_modified_local":dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")})
    total = len(all_accounts)
    identical_counts = len({m["accounts"] for m in monthly}) == 1 and len(monthly) >= 3
    if identical_counts:
        issues.append(f"모든 {len(monthly)}개 파일의 고유 계정 수가 {monthly[0]['accounts']}개로 동일. 내보내기 건수 제한 및 전체성 확인 필요.")
    coverage = {m["month"] for m in monthly}
    start_date = dt.date.fromisoformat(monthly[0]["month"] + "-01")
    last_date = dt.date.fromisoformat(monthly[-1]["month"] + "-01")
    expected = []
    current = start_date
    while current <= last_date:
        expected.append(current.strftime("%Y-%m"))
        current = dt.date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    missing_months = sorted(set(expected) - coverage)
    lookup = {m["month"]: m["accounts"] for m in monthly}
    count_checks = []
    if reference_counts:
        with reference_counts.open(encoding='utf-8-sig',newline='') as handle:
            for ref in csv.DictReader(handle):
                if ref['month'] not in lookup:
                    continue
                expected_count = int(ref['total_accounts'])
                if expected_count <= 0:
                    raise ValueError('Reference account total must be positive')
                observed_count = lookup[ref['month']]
                approximate = ref.get('approximate','false').lower() == 'true'
                mismatch = observed_count != expected_count
                count_checks.append({'month':ref['month'],'reported_total_accounts':expected_count,
                                     'reported_total_is_approximate':approximate,'observed_accounts':observed_count,
                                     'observed_fraction':observed_count/expected_count,'counts_differ':mismatch,
                                     'source':ref.get('source','')})
                if mismatch:
                    qualifier='약 ' if approximate else ''
                    issues.append(f"{ref['month']} 조회 건수 {qualifier}{expected_count:,}개와 파일 {observed_count:,}개가 다릅니다. 비율은 제공분 포함률이며 통계적 표본 추출률이 아닙니다.")
    growth_ready = completeness_confirmed and not missing_months and not any(c['counts_differ'] for c in count_checks)
    for i, m in enumerate(monthly):
        year, mon = map(int, m["month"].split("-"))
        prev = f"{year - (mon == 1):04}-{12 if mon == 1 else mon-1:02}"
        m["mom_growth"] = rates(m["accounts"], lookup.get(prev)) if growth_ready else None
        m["yoy_growth"] = rates(m["accounts"], lookup.get(f"{year-1:04}-{mon:02}")) if growth_ready else None
    annual = []
    matched_periods = []
    years = sorted({int(m["month"][:4]) for m in monthly})
    for year in years:
        available = [m for m in monthly if m["month"].startswith(str(year))]
        full = len(available) == 12
        prior = [m for m in monthly if m["month"].startswith(str(year - 1))]
        total_year = sum(m["accounts"] for m in available)
        annual.append({"year":year,"months_available":len(available),"observed_accounts":total_year,
                       "full_year_growth":rates(total_year,sum(m["accounts"] for m in prior)) if growth_ready and full and len(prior)==12 else None})
        last_month = max(int(m['month'][5:]) for m in available)
        wanted_current = [f'{year:04}-{month:02}' for month in range(1,last_month+1)]
        wanted_previous = [f'{year-1:04}-{month:02}' for month in range(1,last_month+1)]
        has_current = all(month in lookup for month in wanted_current)
        has_previous = all(month in lookup for month in wanted_previous)
        current_count = sum(lookup[month] for month in wanted_current) if has_current else None
        previous_count = sum(lookup[month] for month in wanted_previous) if has_previous else None
        matched_periods.append({'year':year,'through_month':last_month,'current_observed_accounts':current_count,
                                'previous_observed_accounts':previous_count,
                                'same_period_growth':rates(current_count,previous_count) if growth_ready and has_current and has_previous else None})
    def totals(counter, names):
        return [{**dict(zip(names,k if isinstance(k,tuple) else (k,))),"accounts":v,"sample_share":v/total} for k,v in sorted(counter.items(), key=lambda kv:(-kv[1],str(kv[0])))]
    result={"scope":{"first_month":monthly[0]["month"],"last_month":monthly[-1]["month"],"file_count":len(files),
                     "registration_filter":"User confirmed: month start 00:00 through month end 23:59. Filter timezone not confirmed.",
                     "account_unit":"Distinct numeric account IDs, not distinct people",
                     "language_basis":"Current app setting at export; no language history",
                     "completeness_confirmed":completeness_confirmed,"growth_available":growth_ready,
                     "blank_counts_are_zero":blanks_are_zero,"missing_months":missing_months,
                     "created_at":dt.datetime.now().astimezone().isoformat(timespec="seconds")},
            "totals":{"raw_rows":sum(m['raw_rows'] for m in monthly),"accounts":total,
                      "duplicate_rows_removed":sum(m['duplicate_rows_removed'] for m in monthly),
                      "cross_month_duplicate_accounts":0,
                      **{k:sum(m[k] for m in monthly) for k in monthly[0] if k.endswith('_accounts') or k=='recorded_cumulative_orders'}},
            "issues":issues,"count_checks":count_checks,"monthly":monthly,"annual":annual,"matched_periods":matched_periods,"languages":totals(language,['language']),
            "countries":totals(country,['country']),"cities":totals(city,['country','city']),
            "country_language":totals(country_language,['country','language']),"levels":totals(level,['level']),
            "monthly_language":[{"month":m,"language":l,"accounts":n,"within_month_share":n/lookup[m]} for (m,l),n in sorted(monthly_language.items())],
            "monthly_country":[{"month":m,"country":c,"accounts":n,"within_month_share":n/lookup[m]} for (m,c),n in sorted(monthly_country.items())],
            "order_buckets":[{"month":m,"bucket":b,"accounts":n} for (m,b),n in sorted(buckets.items())],"sources":manifests}
    for name in ['languages','countries','cities','country_language','levels','monthly_language','monthly_country','order_buckets']:
        assert sum(r['accounts'] for r in result[name]) == total, name
    assert result['totals']['raw_rows']-result['totals']['duplicate_rows_removed']==total
    output_dir.mkdir(parents=True,exist_ok=True)
    (output_dir/'analysis.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    for name in ['monthly','annual','matched_periods','count_checks','languages','countries','cities','country_language','monthly_language','monthly_country','order_buckets','sources']:
        csv_out(output_dir/f'{name}.csv',result[name])
    print(json.dumps({"scope":result['scope'],"totals":result['totals'],"issues":issues,"languages":result['languages'],"countries":result['countries'][:12],"levels":result['levels']},ensure_ascii=False))
    return result


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--input',type=pathlib.Path,required=True)
    parser.add_argument('--output',type=pathlib.Path,required=True)
    parser.add_argument('--start')
    parser.add_argument('--end')
    parser.add_argument('--completeness-confirmed',action='store_true')
    parser.add_argument('--blank-counts-are-zero',action='store_true')
    parser.add_argument('--reference-counts',type=pathlib.Path)
    args=parser.parse_args()
    analyze(args.input,args.output,args.start,args.end,args.completeness_confirmed,args.blank_counts_are_zero,args.reference_counts)
