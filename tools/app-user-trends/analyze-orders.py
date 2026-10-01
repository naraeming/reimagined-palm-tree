"""Read aggregate order exports without treating their rows as individual orders."""
import argparse
import csv
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re
import unicodedata

FIELDS=['currency','country','city','industry','language','status','product_amount','final_amount','orders']
DEFAULT_STATUS={
    'blank':{'label':'주문취소','group':'cancelled'},
    '1':{'label':'주문접수','group':'known_other'},
    '2':{'label':'주문확인','group':'known_other'},
    '3':{'label':'조리중','group':'known_other'},
    '4':{'label':'배달시작','group':'known_other'},
    '6':{'label':'배달완료','group':'completed'},
}

def num(value):
    if value is None or isinstance(value,bool):
        raise ValueError('Missing numeric value')
    try:
        result=Decimal(str(value).replace(',','').strip())
    except InvalidOperation as exc:
        raise ValueError('Invalid numeric value') from exc
    if not result.is_finite():
        raise ValueError('Non-finite numeric value')
    return result

def clean(value,upper=False):
    if value is None or not str(value).strip():
        return '미기록'
    value=unicodedata.normalize('NFC',str(value).strip())
    return value.upper() if upper else value.lower()

def convert(value):
    if isinstance(value,Decimal):
        return str(value)
    raise TypeError(type(value).__name__)

def write_csv(path,rows):
    if not rows:return
    with path.open('w',encoding='utf-8-sig',newline='') as handle:
        writer=csv.DictWriter(handle,fieldnames=list(rows[0]))
        writer.writeheader();writer.writerows(rows)

def analyze(paths,output,status_map=None,amounts_are_sums=False,full_month_confirmed=False,control_total=None):
    from openpyxl import load_workbook
    statuses=status_map or DEFAULT_STATUS
    if any(meta.get('group') not in {'cancelled','completed','known_other','unknown'} or not meta.get('label') for meta in statuses.values()):
        raise ValueError('Invalid status mapping')
    normalized=[];sources=[];keys=set();file_months=set()
    for path in sorted(paths):
        if not re.fullmatch(r'\d{4}',path.stem):raise ValueError('Expected YYMM.xlsx filename')
        month=f'20{path.stem[:2]}-{path.stem[2:]}'
        dt.date.fromisoformat(month+'-01')
        if month in file_months:raise ValueError('Multiple files for the same month require explicit partition review')
        file_months.add(month)
        wb=load_workbook(path,read_only=True,data_only=True)
        if len(wb.worksheets)!=1:raise ValueError('Expected one sheet per monthly export')
        data_rows=0
        for row_no,row in enumerate(wb.active.iter_rows(values_only=True),1):
            if all(v is None for v in row):continue
            if len(row)!=9:raise ValueError(f'{path.name}:{row_no}: expected nine columns')
            # Input has no header. The first row is data and is never discarded.
            if not isinstance(row[0],str) or not re.fullmatch('[A-Z]{3}',row[0].strip()):
                raise ValueError(f'{path.name}:{row_no}: currency must be a three-letter code; review header/total rows')
            record=dict(zip(FIELDS,row))
            for key in FIELDS[:5]:record[key]=clean(record[key],key in ('currency','country'))
            status='blank' if row[5] is None or str(row[5]).strip()=='' else str(int(num(row[5])))
            if status!='blank' and num(row[5])!=int(num(row[5])):raise ValueError('Non-integer status')
            meta=statuses.get(status,{'label':f'상태 {status}(정의 미확인)','group':'unknown'})
            record.update(month=month,status=status,status_label=meta['label'],status_group=meta['group'],source_file=path.name,source_row=row_no)
            for key in ('orders','product_amount','final_amount'):record[key]=num(record[key])
            if record['orders']<=0 or record['orders']!=int(record['orders']):raise ValueError('Order count must be a positive integer')
            record['orders']=int(record['orders'])
            key=(month,*(record[k] for k in FIELDS[:6]))
            if key in keys:raise ValueError('Repeated aggregate dimension key: review export before combining')
            keys.add(key)
            normalized.append(record);data_rows+=1
        sources.append({'file':path.name,'month':month,'rows':data_rows,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'header_present':False})
        wb.close()

    def rollup(dimensions,amounts=False):
        groups={}
        for row in normalized:
            key=tuple(row[k] for k in dimensions)
            if key not in groups:
                groups[key]={**dict(zip(dimensions,key)),'orders':0,'completed_orders':0,'cancelled_orders':0,'known_other_orders':0,'unknown_status_orders':0}
                if amounts:
                    groups[key].update(product_amount=Decimal(0),final_amount=Decimal(0),completed_product_amount=Decimal(0),completed_final_amount=Decimal(0))
            target=groups[key];target['orders']+=row['orders']
            group_field={'completed':'completed_orders','cancelled':'cancelled_orders','known_other':'known_other_orders','unknown':'unknown_status_orders'}[row['status_group']]
            target[group_field]+=row['orders']
            if amounts:
                for measure in ('product_amount','final_amount'):
                    target[measure]+=row[measure]
                    if row['status_group']=='completed':target['completed_'+measure]+=row[measure]
        for row in groups.values():
            row['completion_rate']=row['completed_orders']/row['orders']
            row['cancellation_rate']=row['cancelled_orders']/row['orders']
            if amounts:
                row['completed_final_amount_per_order']=row['completed_final_amount']/row['completed_orders'] if amounts_are_sums and row['completed_orders'] else None
        return sorted(groups.values(),key=lambda r:(-r['orders'],str(r)))

    if not normalized:raise ValueError('No aggregate data rows found')
    total=rollup([])[0]
    if control_total is not None and (control_total < 0 or len(file_months)!=1):
        raise ValueError('Control total must be non-negative and used with one month')
    result={'scope':{'months':sorted(file_months),'aggregate_rows':len(normalized),'amounts_are_sums_confirmed':amounts_are_sums,
                     'full_month_confirmed':full_month_confirmed,'admin_control_total':control_total,
                     'admin_total_matches':None if control_total is None else control_total==total['orders'],
                     'language_basis':'주문언어(사용자 지정)',
                     'status_mapping':statuses,'time_basis':'파일명 기준 주문월; 추출 필터 날짜 기준 및 시간대 추가 확인',
                     'raw_member_identifiers_present':False},'totals':total,'sources':sources,'rows':normalized,
            'monthly':rollup(['month']),'status':rollup(['status','status_label','status_group']),
            'countries':rollup(['country']),'cities':rollup(['country','city']),
            'languages':rollup(['language']),'industries':rollup(['industry']),
            'country_language':rollup(['country','language']),'country_industry':rollup(['country','industry']),
            'currency':rollup(['currency'],True),
            'monthly_country':rollup(['month','country']),
            'monthly_language':rollup(['month','language'])}
    # A count must reconcile across every partition, regardless of currency or status.
    for key in ['monthly','status','countries','cities','languages','industries','country_language','country_industry','currency']:
        assert sum(r['orders'] for r in result[key])==total['orders'],key
        assert sum(r['completed_orders'] for r in result[key])==total['completed_orders'],key
        assert sum(r['cancelled_orders'] for r in result[key])==total['cancelled_orders'],key
    assert sum(total[k] for k in ['completed_orders','cancelled_orders','known_other_orders','unknown_status_orders'])==total['orders']
    output.mkdir(parents=True,exist_ok=True)
    (output/'order-analysis.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=convert),encoding='utf-8')
    for key in ['monthly','status','countries','cities','languages','industries','country_language','country_industry','currency','monthly_country','monthly_language','sources']:
        write_csv(output/f'{key}.csv',result[key])
    print(json.dumps({key:result[key] for key in ['scope','totals','status','countries','languages','industries','currency']},ensure_ascii=False,default=convert))
    return result

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--input',nargs='+',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    p.add_argument('--status-map',type=Path);p.add_argument('--amounts-are-sums',action='store_true');p.add_argument('--full-month-confirmed',action='store_true');p.add_argument('--control-total',type=int)
    a=p.parse_args();statuses=json.loads(a.status_map.read_text(encoding='utf-8')) if a.status_map else None
    analyze(a.input,a.output,statuses,a.amounts_are_sums,a.full_month_confirmed,a.control_total)
